#!/usr/bin/env python3
"""
build_excel.py — Build a formatted Excel lead list from one or more cleaned CSVs.
Produces a professional Excel file matching the Fulfillment House lead list standard.
Usage: python build_excel.py <input1.csv> [input2.csv ...] <output.xlsx>
"""

import csv
import sys
import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import AutoFilter


NAVY = "1F3864"
WHITE = "FFFFFF"
GREEN = "C6EFCE"
YELLOW = "FFEB9C"
RED = "FFC7CE"
LIGHT_GRAY = "F2F2F2"
QUALIFIED_BLUE = "DDEEFF"


def make_border():
    thin = Side(style="thin", color="CCCCCC")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def score_lead(lead):
    """Score a lead 0-5 for qualification. 3+ = qualified."""
    score = 0
    if lead.get("Website", "").strip():
        score += 1
    if lead.get("Email", "").strip():
        score += 2
    if lead.get("Phone", "").strip():
        score += 1
    try:
        rating = float(lead.get("Google Rating", "") or 0)
        if rating >= 4.0:
            score += 1
    except (ValueError, TypeError):
        pass
    return score


def build_excel(csv_files, output_path):
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Lead List"

    headers = ["#", "Business Name", "Website", "Email", "Phone",
               "Address", "City", "State", "Zip",
               "Google Rating", "Review Count", "Category", "Source", "Notes", "Score", "Qualified"]

    col_widths = [5, 35, 35, 32, 18, 30, 20, 8, 12, 12, 12, 22, 18, 30, 8, 10]

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(name="Arial", bold=True, color=WHITE, size=10)
        cell.fill = PatternFill("solid", start_color=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        cell.border = make_border()
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"

    all_leads = []
    for csv_file in csv_files:
        with open(csv_file, "r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for row in reader:
                all_leads.append(row)

    # Score and sort: qualified leads first, then by score desc
    for lead in all_leads:
        lead["_score"] = score_lead(lead)
        lead["_qualified"] = lead["_score"] >= 3

    all_leads.sort(key=lambda x: (-int(x["_qualified"]), -x["_score"]))

    print(f"Total leads loaded: {len(all_leads)}")
    qualified_count = sum(1 for l in all_leads if l["_qualified"])
    print(f"Qualified leads (score >= 3): {qualified_count}")

    rating_col = headers.index("Google Rating") + 1

    for row_idx, lead in enumerate(all_leads, 2):
        is_qualified = lead["_qualified"]
        row_fill = PatternFill("solid", start_color=QUALIFIED_BLUE) if is_qualified else (
            PatternFill("solid", start_color=LIGHT_GRAY) if row_idx % 2 == 0 else None
        )

        for col_idx, header in enumerate(headers, 1):
            if header == "#":
                value = row_idx - 1
            elif header == "Score":
                value = lead["_score"]
            elif header == "Qualified":
                value = "Yes" if is_qualified else ""
            else:
                value = lead.get(header, "")
                if header == "Google Rating" and value:
                    try:
                        value = float(value)
                    except ValueError:
                        pass

            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = Font(name="Arial", size=9, bold=True if is_qualified and col_idx == 2 else False)
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = make_border()

            if row_fill:
                cell.fill = row_fill

        rating_cell = ws.cell(row=row_idx, column=rating_col)
        try:
            rating = float(lead.get("Google Rating", "") or 0)
            if rating >= 4.5:
                rating_cell.fill = PatternFill("solid", start_color=GREEN)
            elif rating >= 3.5:
                rating_cell.fill = PatternFill("solid", start_color=YELLOW)
            elif 0 < rating < 3.5:
                rating_cell.fill = PatternFill("solid", start_color=RED)
        except (ValueError, TypeError):
            pass

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(all_leads) + 1}"

    # ── Summary sheet ────────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Summary")

    def s_header(row, label):
        cell = ws2.cell(row=row, column=1, value=label)
        cell.font = Font(name="Arial", bold=True, color=WHITE, size=10)
        cell.fill = PatternFill("solid", start_color=NAVY)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        ws2.merge_cells(f"A{row}:B{row}")

    def s_row(row, label, formula):
        label_cell = ws2.cell(row=row, column=1, value=label)
        label_cell.font = Font(name="Arial", size=9)
        label_cell.alignment = Alignment(horizontal="left", vertical="center")
        label_cell.border = make_border()
        val_cell = ws2.cell(row=row, column=2, value=formula)
        val_cell.font = Font(name="Arial", bold=True, size=9)
        val_cell.alignment = Alignment(horizontal="center", vertical="center")
        val_cell.border = make_border()

    last_data_row = len(all_leads) + 1
    lead_sheet = "'Lead List'"

    ws2.cell(row=1, column=1, value="3PL Lead List — Summary").font = Font(name="Arial", bold=True, size=14, color=NAVY)
    ws2.cell(row=2, column=1, value=f"Generated: {datetime.now().strftime('%B %d, %Y')}").font = Font(name="Arial", size=9, italic=True)

    s_header(4, "Lead Counts")
    s_row(5, "Total Leads", f"=COUNTA({lead_sheet}!B2:B{last_data_row})")
    s_row(6, "Qualified Leads (score 3+)", f"=COUNTIF({lead_sheet}!P2:P{last_data_row},\"Yes\")")
    s_row(7, "Leads with Email", f"=COUNTIF({lead_sheet}!D2:D{last_data_row},\"?*\")")
    s_row(8, "Leads with Phone", f"=COUNTIF({lead_sheet}!E2:E{last_data_row},\"?*\")")
    s_row(9, "Leads with Website", f"=COUNTIF({lead_sheet}!C2:C{last_data_row},\"?*\")")

    s_header(11, "Ratings")
    s_row(12, "Average Rating", f"=IFERROR(AVERAGEIF({lead_sheet}!J2:J{last_data_row},\">=1\"),\"—\")")
    s_row(13, "Rated 4.5+", f"=COUNTIF({lead_sheet}!J2:J{last_data_row},\">=4.5\")")
    s_row(14, "Rated 3.5-4.4", f"=COUNTIFS({lead_sheet}!J2:J{last_data_row},\">=3.5\",{lead_sheet}!J2:J{last_data_row},\"<4.5\")")
    s_row(15, "Rated below 3.5", f"=COUNTIFS({lead_sheet}!J2:J{last_data_row},\">0\",{lead_sheet}!J2:J{last_data_row},\"<3.5\")")

    s_header(17, "By Source")
    sources = list(set(l.get("Source", "Unknown") for l in all_leads))
    for i, source in enumerate(sorted(sources)):
        s_row(18 + i, source, f"=COUNTIF({lead_sheet}!M2:M{last_data_row},\"{source}\")")

    ws2.column_dimensions["A"].width = 30
    ws2.column_dimensions["B"].width = 16

    os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else ".", exist_ok=True)
    wb.save(output_path)
    print(f"Excel saved: {output_path}")
    print(f"   Leads: {len(all_leads)} | Qualified: {qualified_count} | Sheets: Lead List + Summary")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python build_excel.py <input1.csv> [input2.csv ...] <output.xlsx>")
        sys.exit(1)

    *inputs, output = sys.argv[1:]
    build_excel(inputs, output)
